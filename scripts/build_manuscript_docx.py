#!/usr/bin/env python3
"""Build the anonymous journal manuscript as a publication-ready DOCX."""

from __future__ import annotations

import csv
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript.md"
RESULTS = ROOT / "results"
FIGURES = RESULTS / "figures"
WORK = ROOT / "work"
INSERT_NOTE = Path(
    r"C:\Users\shi16\.codex\plugins\cache\openai-primary-runtime\documents\26.819.11345"
    r"\skills\documents\scripts\insert_note.py"
)

ACCOUNT_ORDER = ["dnc", "harris", "rnc", "trump"]
ACCOUNT_LABEL = {"dnc": "DNC", "harris": "Harris", "rnc": "RNC", "trump": "Trump"}


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5, bold=None, italic=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="666666", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, size=9)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(21)

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True

    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(18)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.55)
    section.right_margin = Cm(2.55)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.1)


INLINE_RE = re.compile(r"(\*\*.*?\*\*|\*.*?\*|\[\[FN\d{3}\]\])")


def add_inline(paragraph, text, default_size=10.5):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = not bold and part.startswith("*") and part.endswith("*")
        content = part[2:-2] if bold else part[1:-1] if italic else part
        run = paragraph.add_run(content)
        set_run_font(run, size=default_size, bold=bold if bold else None, italic=italic if italic else None)


def add_table_title(doc, number, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"表{number}  {title}")
    set_run_font(run, east_asia="黑体", size=10.5, bold=True)


def add_table_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("注：" + text)
    set_run_font(r, size=8.5)


def format_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            if widths:
                cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri == 0:
                set_cell_shading(cell, "E7E7E7")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=8.5, bold=(ri == 0))


def add_table1(doc):
    rows = {}
    with (RESULTS / "table1_descriptive.csv").open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            rows[row["account_slug"]] = row
    add_table_title(doc, 1, "四个官方账号的样本与累计平台指标")
    headers = ["账号", "角色", "样本量", "日均发布", "时长中位数（秒）", "观看中位数", "点赞率中位数", "评论率中位数"]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for slug in ACCOUNT_ORDER:
        x = rows[slug]
        vals = [
            ACCOUNT_LABEL[slug],
            "党组织" if x["actor_type"] == "party" else "候选人",
            x["n"],
            f'{float(x["posts_per_day"]):.2f}',
            f'{float(x["duration_median"]):.0f}',
            f'{float(x["views_median"]):,.0f}',
            f'{float(x["like_rate_median"])*100:.2f}%',
            f'{float(x["comment_rate_median"])*100:.2f}%',
        ]
        cells = table.add_row().cells
        for cell, value in zip(cells, vals):
            cell.text = value
    format_table(table, [1.4, 1.5, 1.2, 1.5, 2.0, 2.1, 2.1, 2.1])
    add_table_note(doc, "观看、点赞和评论为2026年8月24日抓取时点累计值；点赞率和评论率分别以观看数为分母。")


def add_table2(doc):
    wanted = [
        ("opponent_reference", "对手指涉"),
        ("attack_signal", "攻击"),
        ("mobilization_signal", "动员"),
        ("policy_signal", "政策"),
        ("personal_signal", "私人／日常"),
        ("endorsement_signal", "背书"),
        ("first_person_signal", "第一人称"),
    ]
    values = {}
    with (RESULTS / "table2_text_signals.csv").open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            values[(row["account_slug"], row["signal"])] = float(row["proportion"]) * 100
    add_table_title(doc, 2, "四个官方账号标题信号比例（%）")
    headers = ["标题信号", "DNC", "Harris", "RNC", "Trump"]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for key, label in wanted:
        cells = table.add_row().cells
        vals = [label] + [f"{values[(slug, key)]:.1f}" for slug in ACCOUNT_ORDER]
        for cell, value in zip(cells, vals):
            cell.text = value
    format_table(table, [3.1, 2.8, 2.8, 2.8, 2.8])
    add_table_note(doc, "同一标题可同时命中多个信号。词典规则见随稿数据说明；比例不代表完整视频语义编码。")


def add_figure(doc, number, filename, caption, note):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURES / filename), width=Cm(15.6))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(1)
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(f"图{number}  {caption}")
    set_run_font(r, east_asia="黑体", size=9.5, bold=True)
    add_table_note(doc, note)


def add_body_paragraph(doc, text, kind="body"):
    p = doc.add_paragraph()
    if kind == "abstract":
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(10.5)
        p.paragraph_format.right_indent = Pt(10.5)
        p.paragraph_format.space_after = Pt(3)
    elif kind == "keywords":
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(10.5)
        p.paragraph_format.right_indent = Pt(10.5)
        p.paragraph_format.space_after = Pt(6)
    elif kind == "rq":
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(21)
        p.paragraph_format.space_after = Pt(2)
    elif kind == "reference":
        p.paragraph_format.first_line_indent = Pt(-18)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(15)
        p.paragraph_format.space_after = Pt(2)
    elif kind == "english":
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(16)
        p.paragraph_format.space_after = Pt(4)
    add_inline(p, text, default_size=9.5 if kind in {"abstract", "keywords", "reference", "english"} else 10.5)
    return p


def parse_source():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    notes = {}
    body = []
    for line in lines:
        m = re.match(r"^\[\[FN(\d{3})\]\]\s+(.*)$", line)
        if m:
            notes[f"[[FN{m.group(1)}]]"] = m.group(2)
        else:
            body.append(line)
    return body, notes


def build_base(out_path):
    doc = Document()
    configure_section(doc.sections[0])
    setup_styles(doc)
    add_page_number(doc.sections[0].footer.paragraphs[0])
    body, notes = parse_source()
    in_references = False
    in_english = False
    first_title = True

    for raw in body:
        line = raw.strip()
        if not line:
            continue
        if line == "[[TABLE1]]":
            add_table1(doc)
            continue
        if line == "[[TABLE2]]":
            add_table2(doc)
            continue
        if line == "[[FIGURE1]]":
            add_figure(doc, 1, "figure1_output_and_attention.png", "样本期发布量与单条累计观看中位数", "右图使用对数坐标；累计观看为2026年8月24日抓取值。")
            continue
        if line == "[[FIGURE2]]":
            add_figure(doc, 2, "figure2_lexical_signal_heatmap.png", "四账号标题中的可复核传播信号", "数值为各账号标题命中固定词典的比例；同一标题可同时命中多个信号。")
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.paragraph_format.first_line_indent = Pt(0)
            add_inline(p, line[2:], default_size=18)
            first_title = False
            continue
        if line.startswith("## ——"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(line[3:])
            set_run_font(r, east_asia="黑体", size=13, bold=True)
            continue
        if line.startswith("## "):
            heading = line[3:]
            in_references = heading == "参考文献"
            in_english = heading == "English Title and Abstract"
            doc.add_paragraph(heading, style="Heading 1")
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            continue
        if line.startswith("**内容提要：**"):
            add_body_paragraph(doc, line, "abstract")
            continue
        if line.startswith("**关键词：**"):
            add_body_paragraph(doc, line, "keywords")
            continue
        if re.match(r"^RQ\d：", line):
            add_body_paragraph(doc, line, "rq")
            continue
        if in_references and re.match(r"^\d+\.\s", line):
            add_body_paragraph(doc, line, "reference")
            continue
        if in_english:
            p = add_body_paragraph(doc, line, "english")
            if line.startswith("**Has the Party"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(11)
            continue
        add_body_paragraph(doc, line)

    props = doc.core_properties
    props.title = "党组织退场了吗？平台化竞选中的“政党—候选人”传播分工"
    props.subject = "匿名投稿稿件"
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    doc.save(out_path)
    return notes


def patch_footnotes(docx_path):
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        with zipfile.ZipFile(docx_path, "r") as zin:
            zin.extractall(tmp)

        from lxml import etree

        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        settings_path = tmp / "word" / "settings.xml"
        tree = etree.parse(str(settings_path))
        root = tree.getroot()
        footnote_pr = root.find("w:footnotePr", ns)
        if footnote_pr is None:
            footnote_pr = etree.Element(f"{{{ns['w']}}}footnotePr")
            root.insert(0, footnote_pr)
        restart = footnote_pr.find("w:numRestart", ns)
        if restart is None:
            restart = etree.SubElement(footnote_pr, f"{{{ns['w']}}}numRestart")
        restart.set(f"{{{ns['w']}}}val", "eachPage")
        tree.write(str(settings_path), xml_declaration=True, encoding="UTF-8", standalone="yes")

        document_path = tmp / "word" / "document.xml"
        dtree = etree.parse(str(document_path))
        droot = dtree.getroot()
        for ref in droot.xpath(".//w:footnoteReference", namespaces=ns):
            run = ref.getparent()
            rpr = run.find("w:rPr", ns)
            if rpr is None:
                rpr = etree.Element(f"{{{ns['w']}}}rPr")
                run.insert(0, rpr)
            vert = rpr.find("w:vertAlign", ns)
            if vert is None:
                vert = etree.SubElement(rpr, f"{{{ns['w']}}}vertAlign")
            vert.set(f"{{{ns['w']}}}val", "superscript")
            sz = rpr.find("w:sz", ns)
            if sz is None:
                sz = etree.SubElement(rpr, f"{{{ns['w']}}}sz")
            sz.set(f"{{{ns['w']}}}val", "16")
        dtree.write(str(document_path), xml_declaration=True, encoding="UTF-8", standalone="yes")

        footnotes_path = tmp / "word" / "footnotes.xml"
        if footnotes_path.exists():
            ftree = etree.parse(str(footnotes_path))
            froot = ftree.getroot()
            for run in froot.xpath(".//w:r", namespaces=ns):
                rpr = run.find("w:rPr", ns)
                if rpr is None:
                    rpr = etree.Element(f"{{{ns['w']}}}rPr")
                    run.insert(0, rpr)
                rfonts = rpr.find("w:rFonts", ns)
                if rfonts is None:
                    rfonts = etree.SubElement(rpr, f"{{{ns['w']}}}rFonts")
                rfonts.set(f"{{{ns['w']}}}ascii", "Times New Roman")
                rfonts.set(f"{{{ns['w']}}}hAnsi", "Times New Roman")
                rfonts.set(f"{{{ns['w']}}}eastAsia", "宋体")
                sz = rpr.find("w:sz", ns)
                if sz is None:
                    sz = etree.SubElement(rpr, f"{{{ns['w']}}}sz")
                sz.set(f"{{{ns['w']}}}val", "18")
                szcs = rpr.find("w:szCs", ns)
                if szcs is None:
                    szcs = etree.SubElement(rpr, f"{{{ns['w']}}}szCs")
                szcs.set(f"{{{ns['w']}}}val", "18")
                if run.find("w:footnoteRef", ns) is not None:
                    vert = rpr.find("w:vertAlign", ns)
                    if vert is None:
                        vert = etree.SubElement(rpr, f"{{{ns['w']}}}vertAlign")
                    vert.set(f"{{{ns['w']}}}val", "superscript")
            ftree.write(str(footnotes_path), xml_declaration=True, encoding="UTF-8", standalone="yes")

        rebuilt = docx_path.with_suffix(".patched.docx")
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as zout:
            for path in tmp.rglob("*"):
                if path.is_file():
                    zout.write(path, path.relative_to(tmp).as_posix())
        rebuilt.replace(docx_path)


def main():
    if len(sys.argv) != 2:
        raise SystemExit("Usage: build_manuscript_docx.py OUTPUT.docx")
    output = Path(sys.argv[1]).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(exist_ok=True)
    base = WORK / "manuscript_base.docx"
    notes = build_base(base)

    current = base
    intermediates = []
    for index, (marker, note_text) in enumerate(sorted(notes.items()), start=1):
        next_path = WORK / f"manuscript_note_{index:02d}.docx"
        subprocess.run(
            [sys.executable, str(INSERT_NOTE), str(current), "--kind", "footnote", "--marker", marker,
             "--text", note_text, "--out", str(next_path)],
            check=True,
        )
        if current != base:
            intermediates.append(current)
        current = next_path
    shutil.copy2(current, output)
    patch_footnotes(output)
    print(f"Created {output}")
    print(f"Footnotes inserted: {len(notes)}")


if __name__ == "__main__":
    main()
